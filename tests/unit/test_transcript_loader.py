import zipfile
from pathlib import Path

import pytest
from docx import Document

from scribe.utils.transcript_loader import (
    _read_docx,
    _read_text_with_fallback,
    _validate_docx_archive,
    is_supported_transcript_file,
    load_transcript_text,
)


def test_is_supported_transcript_file_checks_extensions():
    assert is_supported_transcript_file(Path("meeting.txt")) is True
    assert is_supported_transcript_file(Path("meeting.md")) is True
    assert is_supported_transcript_file(Path("meeting.docx")) is True
    assert is_supported_transcript_file(Path("meeting.pdf")) is False


def test_read_text_with_fallback_handles_utf8_and_utf16(tmp_path):
    utf8_path = tmp_path / "meeting.txt"
    utf8_path.write_text("hello world", encoding="utf-8")
    assert _read_text_with_fallback(utf8_path) == "hello world"

    utf16_path = tmp_path / "meeting-utf16.txt"
    utf16_path.write_text("hello utf16", encoding="utf-16")
    assert _read_text_with_fallback(utf16_path) == "hello utf16"


def test_read_text_with_fallback_rejects_oversized_input(tmp_path, monkeypatch):
    transcript_path = tmp_path / "large.txt"
    transcript_path.write_bytes(b"12345")
    monkeypatch.setattr("scribe.utils.transcript_loader.MAX_TRANSCRIPT_BYTES", 4)
    with pytest.raises(ValueError, match="input limit"):
        _read_text_with_fallback(transcript_path)


def test_read_docx_and_load_transcript_text_support_text_markdown_and_docx(tmp_path):
    docx_path = tmp_path / "meeting.docx"
    document = Document()
    document.add_paragraph("First line")
    document.add_paragraph("Second line")
    document.save(str(docx_path))

    text_path = tmp_path / "meeting.txt"
    text_path.write_text("plain text", encoding="utf-8")
    markdown_path = tmp_path / "meeting.md"
    markdown_path.write_text("# heading", encoding="utf-8")

    assert _read_docx(docx_path) == "First line\nSecond line"
    assert load_transcript_text(text_path) == "plain text"
    assert load_transcript_text(markdown_path) == "# heading"
    assert load_transcript_text(docx_path) == "First line\nSecond line"


def test_read_docx_and_load_transcript_text_raise_clear_errors(tmp_path):
    broken_docx = tmp_path / "broken.docx"
    broken_docx.write_text("not really a docx", encoding="utf-8")

    with pytest.raises(ValueError, match="valid ZIP package"):
        _read_docx(broken_docx)

    with pytest.raises(ValueError, match="Unsupported transcript format"):
        load_transcript_text(tmp_path / "meeting.rtf")


def test_validate_docx_archive_rejects_suspicious_expansion(tmp_path, monkeypatch):
    archive_path = tmp_path / "expanded.docx"
    with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 10_000)

    monkeypatch.setattr("scribe.utils.transcript_loader.MAX_DOCX_ENTRY_BYTES", 100)
    with pytest.raises(ValueError, match="oversized entry"):
        _validate_docx_archive(archive_path)
