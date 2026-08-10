from types import SimpleNamespace
from unittest.mock import Mock, patch

import pytest
from docx import Document

from scribe.exceptions import FileOperationError
from scribe.utils.docx_converter import txt_to_docx


def test_txt_to_docx_rejects_missing_input(tmp_path):
    missing_path = tmp_path / "missing.txt"

    with pytest.raises(FileOperationError, match="Input file does not exist"):
        txt_to_docx(missing_path)


def test_txt_to_docx_converts_transcript_to_document(tmp_path):
    transcript_path = tmp_path / "meeting.txt"
    transcript_path.write_text(
        "\n".join(
            [
                "Meeting: Weekly Sync",
                "2026-05-08",
                "Duration: 5m",
                "--------------------".replace("-", "\u2500"),
                "[00:00] Hello team",
                "[NO AUDIO - silence]",
                "Plain paragraph",
            ]
        ),
        encoding="utf-8",
    )

    output_path = txt_to_docx(transcript_path)

    assert output_path == transcript_path.with_suffix(".docx")
    assert output_path.exists()

    document = Document(str(output_path))
    paragraph_texts = [
        paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text
    ]

    assert paragraph_texts[0] == "Weekly Sync"
    assert paragraph_texts[1] == "2026-05-08"
    assert paragraph_texts[2] == "Duration: 5m"
    assert "[00:00] Hello team" in paragraph_texts
    assert "[NO AUDIO - silence]" in paragraph_texts
    assert "Plain paragraph" in paragraph_texts


def test_txt_to_docx_wraps_document_save_failures(tmp_path):
    transcript_path = tmp_path / "meeting.txt"
    transcript_path.write_text("Transcription: Demo\nBody text", encoding="utf-8")
    mock_document = Mock()
    mock_document.sections = [SimpleNamespace()]
    mock_document.add_paragraph.return_value = Mock(
        alignment=None,
        add_run=Mock(
            return_value=Mock(font=SimpleNamespace(size=None, color=SimpleNamespace(rgb=None)))
        ),
        paragraph_format=SimpleNamespace(line_spacing=None),
    )
    mock_document.save.side_effect = OSError("disk full")

    with patch("scribe.utils.docx_converter.Document", return_value=mock_document):
        with pytest.raises(FileOperationError, match="Failed to convert TXT to DOCX"):
            txt_to_docx(transcript_path, docx_path=tmp_path / "meeting.docx")
